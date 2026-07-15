#!/usr/bin/env python3
"""Auto-generate downloadable materials (PDF + PowerPoint) for the "Infoblatt"
pages of a moodle-mbz-builder course and wire them into the course.

A page opts in with `"material": true`. For each such page this tool:
  * renders a print PDF that mirrors the FULL Infoblatt (header bar, explanatory
    prose, embedded diagram, tables/lists and the Merke box — the same content
    as the page, not just a title + bullets),
  * renders an editable PPTX (title bar + diagram + key points),
  * collects all PDFs and all PPTX into the hidden "Material für Lehrkräfte"
    section as two `folder`s (a PDF set for printing, a PPTX set for editing).
    The PDF is intentionally NOT inserted after the page — the Infoblatt stays
    alone at its spot in the course.

Key points come from `"material_bullets": [...]` if given, else they are
extracted from the page content (list items, else paragraphs).

It writes a NEW course JSON (default `<course>.materialized.json`) that you then
feed to build_mbz.py. The video summary is intentionally never generated — that
stays a manual NotebookLM placeholder.

Requires: svglib, reportlab, pymupdf (fitz), python-pptx, pillow.
    python -m pip install svglib reportlab pymupdf python-pptx pillow

Usage:
    python make_materials.py course.json [-o course.materialized.json]
    python build_mbz.py course.materialized.json -o course.mbz --wwwroot ...
"""
import argparse, copy, html, json, os, re, sys

GREEN_HEX = "#1d6e56"


def _need():
    try:
        import svglib.svglib, reportlab, fitz, pptx, PIL  # noqa
    except Exception as e:
        sys.exit(f"Missing dependency: {e}\nInstall: python -m pip install "
                 f"svglib reportlab pymupdf python-pptx pillow")


def _enable_font_cache():
    """Make svglib parse each TrueType font file at most ONCE per process.

    svglib's ``FontMap.find_font`` drops the weight/style when it falls through
    to ``register_font`` (it calls ``register_font(family, guessed_file)`` with
    the default normal/normal), so a bold/italic lookup is stored under the base
    family name and the weighted key never hits the family cache. Every
    ``svg2rlg`` render therefore re-runs reportlab's ``TTFont`` extraction on the
    same Arial/… files — ~5 s of the material step for a course with several
    diagrams (each SVG is rendered twice: vector into the PDF, rasterised for the
    PPTX).

    Fix: memoise ``register_font`` per process. A repeat call with identical
    arguments returns the first result without re-registering. reportlab keeps
    the TTFont from the first registration — fonts are designed to be registered
    once and reused across documents — so the rendered PDF/PPTX bytes are
    IDENTICAL; only the redundant font-file parse is skipped. Verified against a
    frozen-clock byte manifest of the built .mbz (output unchanged).
    """
    try:
        from svglib import fonts as _sf
    except Exception:  # noqa: BLE001 - optional dependency
        return
    if getattr(_sf.FontMap, "_o2_font_cache", False):
        return
    _orig = _sf.FontMap.register_font

    def _cached(self, font_family, font_path=None, weight="normal",
                style="normal", rlgFontName=None):
        cache = self.__dict__.setdefault("_o2_regcache", {})
        key = (font_family, font_path, weight, style, rlgFontName)
        if key in cache:
            return cache[key]
        res = _orig(self, font_family, font_path, weight, style, rlgFontName)
        cache[key] = res
        return res

    _sf.FontMap.register_font = _cached
    _sf.FontMap._o2_font_cache = True


def slugify(name):
    s = re.sub(r"^infoblatt[:\s-]*", "", name.strip(), flags=re.I)
    s = re.sub(r"[^a-zA-Z0-9]+", "_", s).strip("_").lower()
    return s or "material"


def first_image(page):
    for img in page.get("images", []):
        p = img if isinstance(img, str) else img.get("file")
        if p:
            return p
    return None


def extract_bullets(page):
    if page.get("material_bullets"):
        return page["material_bullets"]
    content = page.get("content", "") or ""
    items = re.findall(r"<li[^>]*>(.*?)</li>", content, flags=re.S | re.I)
    if not items:
        items = re.findall(r"<p[^>]*>(.*?)</p>", content, flags=re.S | re.I)
    out = []
    for it in items:
        txt = html.unescape(re.sub(r"<[^>]+>", "", it)).strip()
        txt = re.sub(r"\s+", " ", txt)
        if txt and len(txt) > 1:
            out.append(txt)
    return out[:6]


def title_of(page):
    return re.sub(r"^Infoblatt[:\s-]*", "", page.get("name", "Infoblatt")).strip()


def svg_to_pdf(svgpath, pdfpath):
    from svglib.svglib import svg2rlg
    from reportlab.graphics import renderPDF
    _enable_font_cache()
    renderPDF.drawToFile(svg2rlg(svgpath), pdfpath)


def svg_to_png(svgpath, pngpath):
    from svglib.svglib import svg2rlg
    from reportlab.graphics import renderPDF
    import fitz
    _enable_font_cache()
    tmp = pngpath + ".tmp.pdf"
    renderPDF.drawToFile(svg2rlg(svgpath), tmp)
    doc = fitz.open(tmp)
    doc[0].get_pixmap(matrix=fitz.Matrix(3, 3), alpha=False).save(pngpath)
    doc.close(); os.remove(tmp)


# ---- minimal HTML -> reportlab flowables (renders the FULL Infoblatt) --------
# The PDF mirrors the page: header bar, explanatory prose, embedded diagram,
# tables/lists and the Merke box — not just a title + bullets.

from html.parser import HTMLParser as _HTMLParser

from fraction_render import to_slash

_BLOCK = {"p", "div", "table", "ul", "ol", "img", "h1", "h2", "h3", "h4", "h5"}
# Glyphs the built-in PDF font (Helvetica/WinAnsi) cannot draw -> strip them so
# they don't become tofu boxes (emoji, dingbats, arrows, variation selectors).
_EMOJI = re.compile(
    "[\U0001F000-\U0001FAFF\U00002600-\U000027BF\U00002B00-\U00002BFF"
    "\U0001F1E6-\U0001F1FF️←-⇿⌀-⏿]")


class _Node:
    __slots__ = ("tag", "attrs", "children", "text")

    def __init__(self, tag, attrs=None):
        self.tag = tag
        self.attrs = dict(attrs or {})
        self.children = []
        self.text = None


class _DOM(_HTMLParser):
    _VOID = {"br", "img", "hr"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = _Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = _Node(tag, attrs)
        self.stack[-1].children.append(node)
        if tag not in self._VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].children.append(_Node(tag, attrs))

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        if data:
            t = _Node("#text")
            t.text = data
            self.stack[-1].children.append(t)


def _parse_html(s):
    p = _DOM()
    p.feed(s or "")
    return p.root


def _clean(t):
    return re.sub(r"\s+", " ", _EMOJI.sub("", (t or "").replace("→", "->")))


def _esc(t):
    return _clean(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _inline(node):
    out = []
    for c in node.children:
        if c.tag == "#text":
            out.append(_esc(c.text))
        elif c.tag in ("b", "strong"):
            out.append("<b>" + _inline(c) + "</b>")
        elif c.tag in ("i", "em"):
            out.append("<i>" + _inline(c) + "</i>")
        elif c.tag == "sub":
            # ReportLab-Paragraph kennt <sub>/<super> nativ -> Indizes wie
            # A<sub>1</sub> erscheinen auch im Druck-PDF tiefgestellt.
            out.append("<sub>" + _inline(c) + "</sub>")
        elif c.tag == "sup":
            out.append("<super>" + _inline(c) + "</super>")
        elif c.tag == "br":
            out.append("<br/>")
        else:
            out.append(_inline(c))
    return "".join(out)


def _style(n):
    return n.attrs.get("style", "")


def _bg(style):
    m = re.search(r"background(?:-color)?\s*:\s*(#[0-9a-fA-F]{3,6})", style)
    return m.group(1) if m else None


def _leftaccent(style):
    m = re.search(r"border-left\s*:[^;]*?(#[0-9a-fA-F]{3,6})", style)
    return m.group(1) if m else None


def _has_block(node):
    return any(c.tag in _BLOCK for c in node.children)


def _find(node, tag):
    out = []
    for c in node.children:
        if c.tag == tag:
            out.append(c)
        out.extend(_find(c, tag))
    return out


def make_pdf(page, base, out):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, ListFlowable, ListItem)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from svglib.svglib import svg2rlg
    _enable_font_cache()

    imgs = {}
    for p in page.get("images", []):
        pth = p if os.path.isabs(p) else os.path.join(base, p)
        imgs[os.path.basename(p)] = pth

    ss = getSampleStyleSheet()
    st = {
        "bar": ParagraphStyle("bar", parent=ss["Normal"], fontSize=13, leading=17,
                              textColor=colors.white, fontName="Helvetica-Bold"),
        "body": ParagraphStyle("body", parent=ss["Normal"], fontSize=10.5, leading=15),
        "h3": ParagraphStyle("h3", parent=ss["Normal"], fontSize=13, leading=16,
                             spaceBefore=8, spaceAfter=3, fontName="Helvetica-Bold"),
        "th": ParagraphStyle("th", parent=ss["Normal"], fontSize=9.5, leading=12,
                             fontName="Helvetica-Bold"),
        "td": ParagraphStyle("td", parent=ss["Normal"], fontSize=9.5, leading=12),
    }
    margin = 16 * mm
    content_w = A4[0] - 2 * margin

    def bar(text, bgcol):
        t = Table([[Paragraph(text, st["bar"])]], colWidths=[content_w])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(bgcol)),
            ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
        return t

    def box(text, bgcol, accent):
        cmds = [("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(bgcol)),
                ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]
        if accent:
            cmds.append(("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor(accent)))
        t = Table([[Paragraph(text, st["body"])]], colWidths=[content_w])
        t.setStyle(TableStyle(cmds))
        return t

    def table(node):
        trs = _find(node, "tr")
        data, spans, ncol, header = [], [], 0, set()
        for ri, tr in enumerate(trs):
            cells = [c for c in tr.children if c.tag in ("td", "th")]
            row, col = [], 0
            for cell in cells:
                cs = int(cell.attrs.get("colspan", 1) or 1)
                sty = st["th"] if cell.tag == "th" else st["td"]
                row.append(Paragraph(_inline(cell).strip() or " ", sty))
                if cell.tag == "th":
                    header.add(ri)
                if cs > 1:
                    spans.append(("SPAN", (col, ri), (col + cs - 1, ri)))
                    row.extend([""] * (cs - 1))
                col += cs
            if _bg(_style(tr)):
                header.add(ri)
            ncol = max(ncol, col)
            data.append(row)
        for r in data:
            r.extend([""] * (ncol - len(r)))
        cw = content_w / max(ncol, 1)
        t = Table(data, colWidths=[cw] * ncol)
        cmds = [("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c9ccd1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
        for ri in header:
            cmds.append(("BACKGROUND", (0, ri), (-1, ri), colors.HexColor("#e1f5ee")))
        cmds.extend(spans)
        t.setStyle(TableStyle(cmds))
        return t

    def image(fn):
        pth = imgs.get(fn)
        if not (pth and os.path.exists(pth)):
            return None
        d = svg2rlg(pth)
        if not d:
            return None
        s = content_w / d.width
        if d.height * s > 150 * mm:
            s = (150 * mm) / d.height
        d.width *= s
        d.height *= s
        d.scale(s, s)
        d.hAlign = "CENTER"
        return d

    story = []

    def render(node):
        for c in node.children:
            tag = c.tag
            if tag == "#text":
                if c.text and c.text.strip():
                    story.append(Paragraph(_esc(c.text), st["body"]))
            elif tag == "div":
                bgcol = _bg(_style(c))
                if _has_block(c):
                    render(c)
                elif bgcol:
                    txt = _inline(c).strip()
                    if not txt:
                        continue
                    if re.search(r"color\s*:\s*(#fff(?:fff)?|white)", _style(c), re.I):
                        story.append(bar(txt, bgcol))
                    else:
                        story.append(box(txt, bgcol, _leftaccent(_style(c))))
                    story.append(Spacer(1, 5))
                else:
                    txt = _inline(c).strip()
                    if txt:
                        story.append(Paragraph(txt, st["body"]))
            elif tag == "p":
                bgcol = _bg(_style(c))
                txt = _inline(c).strip()
                if not txt:
                    continue
                story.append(box(txt, bgcol, _leftaccent(_style(c))) if bgcol
                             else Paragraph(txt, st["body"]))
                story.append(Spacer(1, 4))
            elif tag in ("h2", "h3", "h4"):
                story.append(Paragraph(_inline(c).strip(), st["h3"]))
            elif tag in ("ul", "ol"):
                items = [ci for ci in c.children if ci.tag == "li"]
                flow = [ListItem(Paragraph(_inline(li).strip() or " ", st["body"]),
                                 leftIndent=10) for li in items]
                if flow:
                    story.append(ListFlowable(flow, bulletType="bullet" if tag == "ul" else "1"))
                    story.append(Spacer(1, 5))
            elif tag == "table":
                story.append(table(c))
                story.append(Spacer(1, 5))
            elif tag == "img":
                im = image((c.attrs.get("src", "") or "").split("/")[-1])
                if im is not None:
                    story.append(im)
                    story.append(Spacer(1, 5))
            elif c.children:
                render(c)

    render(_parse_html(to_slash(page.get("content", ""))))
    if not story:
        story = [Paragraph(_esc(title_of(page)), st["h3"])]
    doc = SimpleDocTemplate(out, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm,
                            leftMargin=margin, rightMargin=margin, title=title_of(page))
    doc.build(story)


def make_pptx(title, svgpath, bullets, out, workdir):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.9))
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(0x1d, 0x6e, 0x56); bar.line.fill.background()
    tf = bar.text_frame; tf.margin_left = Inches(0.3); tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = "Infoblatt  ·  " + title
    r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    text_left = 0.5
    if svgpath and os.path.exists(svgpath):
        png = os.path.join(workdir, slugify(title) + ".png")
        svg_to_png(svgpath, png)
        iw, ih = Image.open(png).size; ar = iw / ih
        w = 7.4; hgt = w / ar
        if hgt > 5.4: hgt = 5.4; w = hgt * ar
        slide.shapes.add_picture(png, Inches(0.4), Inches(1.2), width=Inches(w), height=Inches(hgt))
        text_left = 8.1
    box = slide.shapes.add_textbox(Inches(text_left), Inches(1.3), Inches(13.333 - text_left - 0.4), Inches(5.3))
    tf2 = box.text_frame; tf2.word_wrap = True
    for i, b in enumerate(bullets or []):
        para = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        run = para.add_run(); run.text = "•  " + b
        run.font.size = Pt(16); run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
        para.space_after = Pt(10)
    prs.save(out)


def _xesc(s):
    return (str(s or "").replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;"))


def make_glossary_import_xml(gl, out_path):
    """Write a Moodle glossary IMPORT file (Glossar -> Einträge importieren)
    for the given glossary activity. Fallback for the restore quirk that
    glossary entries are user data: restores without 'Nutzerdaten einschließen'
    arrive with an empty glossary — this file re-fills it in two clicks."""
    rows = []
    for e in gl.get("entries", []):
        rows.append(
            "      <ENTRY>\n"
            f"        <CONCEPT>{_xesc(e.get('concept'))}</CONCEPT>\n"
            f"        <DEFINITION>{_xesc(e.get('definition'))}</DEFINITION>\n"
            "        <FORMAT>1</FORMAT>\n"
            "        <USEDYNALINK>1</USEDYNALINK>\n"
            "        <CASESENSITIVE>0</CASESENSITIVE>\n"
            "        <FULLMATCH>1</FULLMATCH>\n"
            "        <TEACHERENTRY>1</TEACHERENTRY>\n"
            "      </ENTRY>"
        )
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        "<GLOSSARY>\n  <INFO>\n"
        f"    <NAME>{_xesc(gl.get('name', 'Glossar'))}</NAME>\n"
        f"    <INTRO>{_xesc(gl.get('intro', ''))}</INTRO>\n"
        "    <INTROFORMAT>1</INTROFORMAT>\n"
        "    <ALLOWDUPLICATEDENTRIES>0</ALLOWDUPLICATEDENTRIES>\n"
        "    <DISPLAYFORMAT>dictionary</DISPLAYFORMAT>\n"
        "    <SHOWSPECIAL>1</SHOWSPECIAL>\n    <SHOWALPHABET>1</SHOWALPHABET>\n"
        "    <SHOWALL>1</SHOWALL>\n    <ALLOWCOMMENTS>0</ALLOWCOMMENTS>\n"
        "    <USEDYNALINK>1</USEDYNALINK>\n    <DEFAULTAPPROVAL>1</DEFAULTAPPROVAL>\n"
        "    <GLOBALGLOSSARY>0</GLOBALGLOSSARY>\n    <ENTBYPAGE>10</ENTBYPAGE>\n"
        "    <ENTRIES>\n" + "\n".join(rows) + "\n    </ENTRIES>\n"
        "  </INFO>\n</GLOSSARY>\n"
    )
    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write(xml)


def ensure_teacher_section(spec):
    for s in spec["sections"]:
        if s.get("visible") == 0:
            return s
    sec = {"name": "Material für Lehrkräfte", "visible": 0,
           "summary": "<p>Nur für Lehrkräfte – Druck- und Editiermaterial.</p>",
           "activities": [{"type": "label", "name": "Hinweis",
                           "intro": "<p>🔒 Dieser Abschnitt ist für Lehrkräfte (verborgen).</p>"}]}
    spec["sections"].append(sec)
    return sec


def main():
    _need()
    ap = argparse.ArgumentParser(description="Auto-generate + wire PDF/PPTX materials for Infoblatt pages.")
    ap.add_argument("course_json")
    ap.add_argument("-o", "--output")
    ap.add_argument("--dir", default="materials", help="output subdir for generated files")
    args = ap.parse_args()

    base = os.path.dirname(os.path.abspath(args.course_json))
    matdir = os.path.join(base, args.dir)
    os.makedirs(matdir, exist_ok=True)
    spec = json.load(open(args.course_json, encoding="utf-8"))

    # Check embedded SVG diagrams BEFORE rasterising them into PDF/PPTX, so a
    # clipped/overlapping diagram is caught before it is baked into materials.
    try:
        import validate_svg
        validate_svg.check_course_svgs(spec, base)
    except Exception as e:  # noqa: BLE001
        print(f"NOTE: SVG diagram check skipped ({e}).", file=sys.stderr)

    pdfs, pptxs = [], []
    for sec in spec["sections"]:
        new_acts = []
        for act in sec["activities"]:
            new_acts.append(act)
            if act.get("type") == "page" and act.get("material"):
                title = title_of(act)
                slug = slugify(act.get("name", "infoblatt"))
                svg = first_image(act)
                bullets = extract_bullets(act)
                pdfp = os.path.join(matdir, slug + ".pdf")
                pptxp = os.path.join(matdir, slug + ".pptx")
                # Attributions-Fuß auch im Druck-PDF (idempotent angehängt)
                pdf_act = act
                try:
                    import build_mbz as _bm
                    pdf_act = dict(act)
                    pdf_act["content"] = _bm.with_attribution(
                        act.get("content", ""), spec)
                except Exception:  # noqa: BLE001
                    pass
                make_pdf(pdf_act, base, pdfp)
                make_pptx(title, svg, bullets, pptxp, matdir)
                pdfs.append(pdfp); pptxs.append(pptxp)
                # The PDF is NOT inserted after the page: the Infoblatt stays
                # alone at its spot in the course; all PDFs are collected in the
                # hidden teacher section below.
        sec["activities"] = new_acts

    if pdfs:
        teacher = ensure_teacher_section(spec)
        teacher["activities"].append({"type": "folder",
            "name": "Alle Infoblätter als PDF (Druckvorlagen)",
            "intro": "<p>Sammlung aller Infoblätter zum Ausdrucken.</p>", "files": pdfs})
        teacher["activities"].append({"type": "folder",
            "name": "Alle Infoblätter als PowerPoint (bearbeitbar)",
            "intro": "<p>Editierbare PPTX-Versionen zum Anpassen.</p>", "files": pptxs})

    # Glossar-Fallback: Einträge sind im Moodle-Backup NUTZERDATEN. Wird beim
    # Wiederherstellen "Nutzerdaten einschließen" abgewählt (oder fehlt das
    # Recht dazu), kommt das Glossar leer an. Für diesen Fall liegt hier eine
    # Import-Datei bereit: Glossar öffnen -> (Zahnrad) Einträge importieren.
    gl_imports = []
    for sec in spec["sections"]:
        for act in sec["activities"]:
            if act.get("type") == "glossary" and act.get("entries"):
                slug = slugify(act.get("name", "glossar"))
                p = os.path.join(matdir, slug + "_import.xml")
                make_glossary_import_xml(act, p)
                gl_imports.append(p)
    if gl_imports:
        teacher = ensure_teacher_section(spec)
        teacher["activities"].append({"type": "resource",
            "name": "Glossar-Importdatei (Fallback)",
            "intro": ("<p><b>Nur nötig, falls das Glossar nach dem Import leer "
                      "ist</b> (das passiert, wenn beim Wiederherstellen "
                      "„Nutzerdaten einschließen“ abgewählt war): Glossar "
                      "öffnen → Zahnrad/„Mehr“ → <b>Einträge importieren</b> → "
                      "diese XML-Datei wählen → importieren. Die Fachbegriffe "
                      "sind danach vollständig da.</p>"),
            "file": gl_imports[0]})

    out = args.output or (os.path.splitext(args.course_json)[0] + ".materialized.json")
    json.dump(spec, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"Materialien erzeugt: {len(pdfs)} PDF + {len(pptxs)} PPTX in {matdir}")
    print(f"Materialisierte Kursdefinition: {out}")


if __name__ == "__main__":
    main()
